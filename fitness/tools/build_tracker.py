#!/usr/bin/env python3
"""Builds fitness/tracker.xlsx (87-day plan: exact meals, training, check-ins).

Usage: python3 fitness/tools/build_tracker.py [--kcal 2300] [--protein 180] [--start 2026-10-07]
The head-coach agent re-runs this with a new --kcal when the weekly check-in changes calories.
"""
import argparse, datetime as dt, itertools, os, random
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ap = argparse.ArgumentParser()
ap.add_argument("--kcal", type=int, default=2300)
ap.add_argument("--protein", type=int, default=175)
ap.add_argument("--start", default="2026-10-07")
ap.add_argument("--days", type=int, default=87)
ap.add_argument("--start-weight", type=float, default=105.0)
ap.add_argument("--target-weight", type=float, default=98.0)
ap.add_argument("--stretch-weight", type=float, default=95.0)
ap.add_argument("--out", default=os.path.join(os.path.dirname(__file__), "..", "tracker.xlsx"))
args = ap.parse_args()
START = dt.date.fromisoformat(args.start)
DAYS = args.days
DIWALI = dt.date(2026, 11, 8)

# ---------------------------------------------------------------- food database
# name: (unit, per, kcal, protein, carbs, fat, fibre)   values are per `per` units
F = {
    "Whey isolate":              ("scoop (30 g)", 1, 110, 25, 1, 0.5, 0),
    "Egg (whole)":               ("pc", 1, 72, 6.3, 0.4, 4.8, 0),
    "Egg white":                 ("pc", 1, 17, 3.6, 0.2, 0, 0),
    "Paneer (full fat)":         ("g", 100, 290, 18, 3, 22, 0),
    "Greek yogurt / hung curd":  ("g", 100, 70, 9, 4, 2, 0),
    "Dahi (toned curd)":         ("g", 100, 60, 3.2, 4.7, 3, 0),
    "Toned milk":                ("ml", 100, 58, 3, 4.7, 3, 0),
    "Soya chunks (dry)":         ("g", 100, 345, 52, 33, 0.5, 13),
    "Tofu":                      ("g", 100, 76, 8, 1.9, 4.8, 0.3),
    "Toor dal (raw wt)":         ("g", 100, 335, 22, 57, 1.5, 15),
    "Moong dal (raw wt)":        ("g", 100, 347, 24, 63, 1.2, 16),
    "Rajma (raw wt)":            ("g", 100, 333, 24, 60, 0.8, 25),
    "Chole (raw wt)":            ("g", 100, 364, 19, 61, 6, 17),
    "Besan":                     ("g", 100, 387, 22, 58, 7, 11),
    "Whole wheat atta (raw wt)": ("g", 100, 340, 12, 71, 1.7, 11),
    "Basmati rice (raw wt)":     ("g", 100, 360, 7, 79, 0.6, 1.3),
    "Rolled oats":               ("g", 100, 389, 17, 66, 7, 10),
    "Banana":                    ("g", 100, 89, 1.1, 23, 0.3, 2.6),
    "Apple":                     ("g", 100, 52, 0.3, 14, 0.2, 2.4),
    "Mixed veg (raw wt)":        ("g", 100, 40, 2, 8, 0.3, 3),
    "Spinach":                   ("g", 100, 23, 2.9, 3.6, 0.4, 2.2),
    "Salad (cucumber/tomato)":   ("g", 100, 15, 0.7, 3.6, 0.1, 0.5),
    "Moong sprouts":             ("g", 100, 30, 3, 6, 0.2, 1.8),
    "Oil / ghee":                ("tsp (5 g)", 1, 45, 0, 0, 5, 0),
    "Almonds":                   ("g", 100, 579, 21, 22, 50, 12.5),
    "Peanut butter":             ("g", 100, 588, 25, 20, 50, 6),
    "Makhana":                   ("g", 100, 347, 9.7, 77, 0.1, 14.5),
    "Flax seeds":                ("g", 100, 534, 18, 29, 42, 27),
    "Roasted chana":             ("g", 100, 370, 22, 58, 5, 17),
    "Mithai (Diwali)":           ("g", 100, 420, 6, 60, 17, 1),
}

def macros(item, qty):
    unit, per, k, p, c, f, fi = F[item]
    m = qty / per
    return [k * m, p * m, c * m, f * m, fi * m]

# ---------------------------------------------------------------- meal templates
# (name, [(item, qty)], flex_item or None)  flex item qty is solved in 10 g (or 1 pc) steps
BREAKFAST = [
    ("Egg bhurji + roti", [("Egg (whole)", 3), ("Egg white", 3), ("Mixed veg (raw wt)", 100), ("Oil / ghee", 1), ("Whole wheat atta (raw wt)", 60)], "Whole wheat atta (raw wt)"),
    ("Protein oats bowl", [("Rolled oats", 50), ("Toned milk", 250), ("Whey isolate", 1), ("Banana", 100), ("Flax seeds", 10)], "Rolled oats"),
    ("Besan chilla stuffed with paneer", [("Besan", 60), ("Paneer (full fat)", 60), ("Mixed veg (raw wt)", 80), ("Oil / ghee", 1), ("Dahi (toned curd)", 100)], "Besan"),
    ("Greek yogurt power bowl", [("Greek yogurt / hung curd", 250), ("Rolled oats", 40), ("Apple", 100), ("Almonds", 10), ("Whey isolate", 1)], "Rolled oats"),
    ("Moong dal chilla + dahi", [("Moong dal (raw wt)", 60), ("Paneer (full fat)", 50), ("Spinach", 50), ("Oil / ghee", 1), ("Dahi (toned curd)", 100)], "Moong dal (raw wt)"),
]
LUNCH = [
    ("Toor dal, rice, paneer sabzi, salad", [("Toor dal (raw wt)", 70), ("Basmati rice (raw wt)", 80), ("Paneer (full fat)", 100), ("Mixed veg (raw wt)", 100), ("Oil / ghee", 1), ("Dahi (toned curd)", 100), ("Salad (cucumber/tomato)", 100)], "Basmati rice (raw wt)"),
    ("Rajma chawal + hung curd", [("Rajma (raw wt)", 70), ("Basmati rice (raw wt)", 70), ("Greek yogurt / hung curd", 150), ("Paneer (full fat)", 50), ("Salad (cucumber/tomato)", 100), ("Oil / ghee", 1)], "Basmati rice (raw wt)"),
    ("Soya chunk curry, roti, dahi", [("Soya chunks (dry)", 50), ("Whole wheat atta (raw wt)", 80), ("Dahi (toned curd)", 150), ("Mixed veg (raw wt)", 100), ("Paneer (full fat)", 50), ("Oil / ghee", 1), ("Salad (cucumber/tomato)", 100)], "Whole wheat atta (raw wt)"),
    ("Chole, rice, hung curd, salad", [("Chole (raw wt)", 70), ("Basmati rice (raw wt)", 70), ("Greek yogurt / hung curd", 150), ("Paneer (full fat)", 50), ("Salad (cucumber/tomato)", 100), ("Oil / ghee", 1)], "Basmati rice (raw wt)"),
    ("Moong dal, roti, egg, sabzi", [("Moong dal (raw wt)", 70), ("Whole wheat atta (raw wt)", 70), ("Egg (whole)", 2), ("Mixed veg (raw wt)", 100), ("Oil / ghee", 1), ("Dahi (toned curd)", 100)], "Whole wheat atta (raw wt)"),
]
SNACK = [  # whey slot, around training
    ("Whey shake + banana + milk", [("Whey isolate", 1), ("Banana", 100), ("Toned milk", 200)]),
    ("Whey shake + Greek yogurt + peanut butter", [("Whey isolate", 1), ("Greek yogurt / hung curd", 200), ("Peanut butter", 12)]),
    ("Whey shake + apple + roasted chana", [("Whey isolate", 1), ("Apple", 150), ("Roasted chana", 30)]),
]
EVENING = [
    ("Roasted makhana + milk", [("Makhana", 25), ("Toned milk", 200)]),
    ("Sprouts chaat with paneer", [("Moong sprouts", 100), ("Paneer (full fat)", 50), ("Salad (cucumber/tomato)", 80)]),
    ("Greek yogurt + almonds", [("Greek yogurt / hung curd", 200), ("Almonds", 10)]),
    ("Roasted chana + dahi", [("Roasted chana", 40), ("Dahi (toned curd)", 100)]),
]
EVENING_DIWALI = ("Diwali: 2 small mithai + Greek yogurt", [("Mithai (Diwali)", 40), ("Greek yogurt / hung curd", 150)])
DINNER = [
    ("Paneer bhurji, roti, sabzi", [("Paneer (full fat)", 150), ("Whole wheat atta (raw wt)", 60), ("Mixed veg (raw wt)", 150), ("Oil / ghee", 1), ("Dahi (toned curd)", 100)], "Whole wheat atta (raw wt)"),
    ("Soya-tofu stir fry + rice", [("Soya chunks (dry)", 40), ("Tofu", 100), ("Basmati rice (raw wt)", 60), ("Mixed veg (raw wt)", 150), ("Oil / ghee", 1)], "Basmati rice (raw wt)"),
    ("Dal, egg curry, roti", [("Toor dal (raw wt)", 60), ("Egg (whole)", 3), ("Whole wheat atta (raw wt)", 60), ("Mixed veg (raw wt)", 150), ("Oil / ghee", 1), ("Salad (cucumber/tomato)", 100)], "Whole wheat atta (raw wt)"),
    ("Palak paneer, roti, dahi", [("Paneer (full fat)", 120), ("Spinach", 150), ("Whole wheat atta (raw wt)", 60), ("Oil / ghee", 1), ("Dahi (toned curd)", 100)], "Whole wheat atta (raw wt)"),
    ("Moong khichdi, paneer, hung curd", [("Moong dal (raw wt)", 40), ("Basmati rice (raw wt)", 40), ("Paneer (full fat)", 100), ("Greek yogurt / hung curd", 150), ("Mixed veg (raw wt)", 150), ("Oil / ghee", 1)], "Basmati rice (raw wt)"),
]

def total(items):
    t = [0.0] * 5
    for it, q in items:
        t = [a + b for a, b in zip(t, macros(it, q))]
    return t

import numpy as np

PROTEIN_ITEMS = ["Paneer (full fat)", "Greek yogurt / hung curd", "Soya chunks (dry)", "Tofu", "Egg (whole)"]

def variants(meal):
    """All (items, macro-vector) variants of a meal: carb/base flex item x protein flex item."""
    flex1 = meal[2] if len(meal) > 2 else None
    flex2 = next((it for it, _ in meal[1] if it in PROTEIN_ITEMS and it != flex1), None)
    d1 = range(-30, 91, 10) if flex1 else [0]
    if flex2 is None:
        d2 = [0]
    elif F[flex2][0] == "pc":
        d2 = [-1, 0, 1, 2]
    else:
        d2 = [-50, -25, 0, 25, 50, 75]
    out = []
    for x, y in itertools.product(d1, d2):
        items, ok = [], True
        for it, q in meal[1]:
            n = q + (x if it == flex1 else y if it == flex2 else 0)
            if (it == flex1 or it == flex2) and not (0.5 * q <= n <= 2.0 * q):
                ok = False
            items.append((it, n))
        if ok:
            out.append((items, np.array(total(items)), abs(x) / 10 + abs(y) / 25))
    return out

_vcache = {}
def V(meal):
    k = meal[0]
    if k not in _vcache:
        _vcache[k] = variants(meal)
    return _vcache[k]

def solve_day(b, l, s, e, d, target, protein, ptol=12, fmax=90):
    base = np.array(total(s[1])) + np.array(total(e[1]))
    vb, vl, vd = V(b), V(l), V(d)
    B = np.array([v[1] for v in vb]); L = np.array([v[1] for v in vl]); D = np.array([v[1] for v in vd])
    pb = np.array([v[2] for v in vb]); pl = np.array([v[2] for v in vl]); pd_ = np.array([v[2] for v in vd])
    tot = base + B[:, None, None, :] + L[None, :, None, :] + D[None, None, :, :]
    pen = pb[:, None, None] + pl[None, :, None] + pd_[None, None, :]
    ok = (np.abs(tot[..., 0] - target) <= 20) & (np.abs(tot[..., 1] - protein) <= ptol) & (tot[..., 3] >= 55) & (tot[..., 3] <= fmax)
    if not ok.any():
        return None
    score = np.abs(tot[..., 0] - target) + 2 * np.abs(tot[..., 1] - protein) + 3 * pen
    score = np.where(ok, score, np.inf)
    i, j, k = np.unravel_index(np.argmin(score), score.shape)
    return (score[i, j, k], vb[i][0], vl[j][0], vd[k][0], list(tot[i, j, k]))

random.seed(7)
combos = list(itertools.product(range(len(BREAKFAST)), range(len(LUNCH)), range(len(SNACK)), range(len(EVENING)), range(len(DINNER))))
random.shuffle(combos)
solved = {}

def get(c, forced_e=None):
    key = (c, forced_e is not None)
    if key not in solved:
        b, l, s, e, d = c
        ev = EVENING_DIWALI if forced_e else EVENING[e]
        solved[key] = solve_day(BREAKFAST[b], LUNCH[l], SNACK[s], ev, DINNER[d], args.kcal, args.protein, *((20, 95) if forced_e else (12, 90)))
    return solved[key]

days = []  # list of dicts
history = []
for n in range(DAYS):
    date = START + dt.timedelta(days=n)
    diwali = date == DIWALI
    chosen = None
    for c in combos[n * 7 % len(combos):] + combos[: n * 7 % len(combos)]:
        # variety: no repeat of the same breakfast/lunch/dinner as either of the previous 2 days
        if any(c[0] == h[0] or c[1] == h[1] or c[4] == h[4] for h in history[-2:]):
            continue
        r = get(c, forced_e=True if diwali else None)
        if r:
            chosen = (c, r)
            break
    if not chosen:
        raise SystemExit(f"No feasible meal combo for {date}; relax constraints")
    c, (score, bi, li, di, tot) = chosen
    history.append(c)
    ev = EVENING_DIWALI if diwali else EVENING[c[3]]
    days.append(dict(date=date, n=n + 1, meals=[
        ("1 Breakfast (~9:00 am)", BREAKFAST[c[0]][0], bi),
        ("2 Lunch (~1:30 pm)", LUNCH[c[1]][0], li),
        ("3 Pre-workout snack (~5:30 pm)", ev[0], ev[1]),
        ("4 Post-workout shake (~9:00 pm)", SNACK[c[2]][0], SNACK[c[2]][1]),
        ("5 Dinner (~9:30 pm)", DINNER[c[4]][0], di),
    ], tot=tot))

# ---------------------------------------------------------------- training
SESSION_BY_WEEKDAY = {0: "Push", 1: "Pull", 2: "Legs", 3: "Rest", 4: "Upper", 5: "Lower", 6: "Rest"}
SESSIONS = {
    "Push": [("Incline DB press", 3, "6-10", "1-2", "Heavy; baseline 30 kg x 8. Add reps to 10, then +2 kg"),
             ("Flat barbell bench press", 3, "6-8", "1-2", "Baseline 60-70 kg x 8-10. Pause 1s on chest"),
             ("Seated DB shoulder press", 2, "8-12", "1-2", ""),
             ("Cable lateral raise", 4, "12-20", "0-1", "Superset with triceps work; priority for aesthetics"),
             ("Triceps pushdown", 3, "10-15", "0-1", ""),
             ("Overhead cable triceps extension", 2, "12-15", "0-1", "")],
    "Pull": [("Lat pulldown (slow 2s eccentric)", 3, "10-15", "1-2", "Stack caps at 67 kg: use tempo/pauses and add reps; single-arm or Smith-machine pull-downs if too light"),
             ("Chest-supported row", 3, "8-12", "1-2", ""),
             ("Seated cable row", 2, "10-12", "1", ""),
             ("Rear delt fly / face pull", 3, "15-20", "0-1", ""),
             ("Incline DB curl", 3, "10-12", "0-1", ""),
             ("Hammer curl", 2, "10-12", "0-1", "")],
    "Legs": [("Barbell squat", 3, "6-10", "1-2", "Baseline 60-70 kg x 8-10"),
             ("Barbell RDL", 3, "8-10", "1-2", "Baseline ~70 kg x 8-10"),
             ("Lying/seated leg curl", 3, "10-15", "0-1", ""),
             ("Standing calf raise", 4, "10-15", "0-1", ""),
             ("Cable crunch", 3, "10-15", "1", "")],
    "Upper": [("Wide-grip pulldown or assisted pull-up", 3, "8-12", "1-2", ""),
              ("Incline machine/DB press", 3, "8-12", "1-2", ""),
              ("Cable lateral raise", 4, "15-20", "0-1", ""),
              ("Cable row (neutral grip)", 3, "10-12", "1", ""),
              ("Reverse pec-deck", 3, "15-20", "0-1", ""),
              ("Cable curl", 3, "12-15", "0-1", ""),
              ("Rope triceps pushdown", 3, "12-15", "0-1", "")],
    "Lower": [("Leg press", 3, "8-12", "1-2", "Baseline 200-240 kg x 8"),
              ("Leg extension", 3, "12-15", "0-1", ""),
              ("Seated leg curl", 2, "10-15", "0-1", ""),
              ("Seated calf raise", 3, "12-15", "0-1", ""),
              ("Hanging leg raise / ab wheel", 3, "10-15", "1", "")],
}
BLOCK = [  # week, phase, set multiplier note, RIR guidance
    (1, "Calibrate", "100% sets", "RIR 2-3; find true working weights. Cap sessions at 75 min"),
    (2, "Build", "100% sets", "RIR 2; add reps first, then load (double progression)"),
    (3, "Build", "100% sets", "RIR 1-2"),
    (4, "Build", "100% sets", "RIR 1-2"),
    (5, "Build", "100% sets", "RIR 1; last compound set to RIR 0-1"),
    (6, "DELOAD", "~60% sets (drop last set of each exercise)", "RIR 3-4; loads -10%; keep steps, keep food"),
    (7, "Build", "100% sets", "RIR 2; resume at pre-deload loads"),
    (8, "Build", "100% sets", "RIR 1-2"),
    (9, "Build", "100% sets", "RIR 1-2"),
    (10, "Build", "100% sets", "RIR 1"),
    (11, "Build", "100% sets", "RIR 0-1 on last set of compounds"),
    (12, "DELOAD", "~60% sets", "RIR 3-4; loads -10%"),
    (13, "Finish", "Light", "Final weigh-in, waist, photos on 1 Jan"),
]

# ---------------------------------------------------------------- workbook
wb = Workbook()
HDR = PatternFill("solid", fgColor="1F3A5F")
SUB = PatternFill("solid", fgColor="DCE6F1")
INP = PatternFill("solid", fgColor="FFF2CC")
def hdr(ws, row, headers, widths=None):
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF"); c.fill = HDR
        c.alignment = Alignment(wrap_text=True, vertical="center")
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

# --- Dashboard
ds = wb.active; ds.title = "Dashboard"
ds["B1"] = "90-day transformation tracker"; ds["B1"].font = Font(bold=True, size=16)
rows = [
    ("Start date", START), ("Start weight (kg)", args.start_weight), ("Target weight (kg) — base case", args.target_weight),
    ("Stretch weight (kg)", args.stretch_weight), ("Deadline", START + dt.timedelta(days=DAYS - 1)),
    ("Daily calories (kcal)", args.kcal), ("Daily protein (g)", args.protein),
]
for i, (k, v) in enumerate(rows, 3):
    ds.cell(row=i, column=2, value=k).font = Font(bold=True)
    c = ds.cell(row=i, column=3, value=v); c.fill = INP
    if isinstance(v, dt.date): c.number_format = "dd-mmm-yyyy"
# C3 start date, C4 start wt, C5 target, C6 stretch, C7 deadline, C8 kcal, C9 protein
ds["B11"] = "Latest weight (kg)"; ds["C11"] = "=IFERROR(LOOKUP(2,1/('Daily Summary'!L2:L200<>\"\"),'Daily Summary'!L2:L200),\"\")"
ds["B12"] = "Latest 7-day avg (kg)"; ds["C12"] = "=IFERROR(LOOKUP(2,1/('Daily Summary'!M2:M200<>\"\"),'Daily Summary'!M2:M200),\"\")"
ds["B13"] = "Lost so far (kg)"; ds["C13"] = '=IF(C12="","",C4-C12)'
ds["B14"] = "Remaining to target (kg)"; ds["C14"] = '=IF(C12="","",C12-C5)'
ds["B15"] = "Days left"; ds["C15"] = "=MAX(0,C7-TODAY())"
ds["B16"] = "Needed pace (kg/week)"; ds["C16"] = '=IF(OR(C12="",C15=0),"",C14/(C15/7))'
ds["B17"] = "Safe pace band (kg/week)"; ds["C17"] = "0.5 - 0.9"
for r in range(11, 18):
    ds.cell(row=r, column=2).font = Font(bold=True)
for r in (11, 12, 13, 14, 16): ds.cell(row=r, column=3).number_format = "0.0"
ds["B19"] = "Yellow cells are inputs. Daily inputs live in 'Daily Summary'; weekly review in 'Weekly Check-in'."
ds.column_dimensions["B"].width = 34; ds.column_dimensions["C"].width = 16

# --- Daily Meals
dm = wb.create_sheet("Daily Meals")
hdr(dm, 1, ["Date", "Day #", "Weekday", "Meal", "Dish", "Item", "Qty", "Unit", "Kcal", "Protein g", "Carbs g", "Fat g", "Fibre g", "Eaten as planned (Y/N)"],
    [12, 7, 11, 24, 36, 28, 8, 12, 8, 10, 9, 8, 9, 14])
r = 2
for d in days:
    for meal, dish, items in d["meals"]:
        for it, q in items:
            k, p, c, f, fi = macros(it, q)
            vals = [d["date"], d["n"], d["date"].strftime("%a"), meal, dish, it, q, F[it][0], round(k), round(p, 1), round(c, 1), round(f, 1), round(fi, 1), ""]
            for ci, v in enumerate(vals, 1):
                cell = dm.cell(row=r, column=ci, value=v)
                if ci == 1: cell.number_format = "dd-mmm-yy"
            r += 1
dm.freeze_panes = "A2"; dm.auto_filter.ref = f"A1:N{r - 1}"
LAST_MEAL_ROW = r - 1

# --- Daily Summary
sm = wb.create_sheet("Daily Summary")
hdr(sm, 1, ["Date", "Day #", "Weekday", "Session", "Trajectory wt (kg)", "Kcal target", "Kcal planned", "Protein g", "Carbs g", "Fat g", "Fibre g",
            "Morning weight (kg)", "7-day avg wt", "Steps", "Sleep (h)", "Readiness 1-5", "Trained? (Y/N)", "Kcal actual", "Notes"],
    [12, 6, 10, 9, 12, 10, 10, 9, 9, 8, 8, 13, 11, 9, 9, 11, 11, 10, 40])
for i, d in enumerate(days, 2):
    wd = d["date"].weekday()
    sm.cell(row=i, column=1, value=d["date"]).number_format = "dd-mmm-yy"
    sm.cell(row=i, column=2, value=d["n"])
    sm.cell(row=i, column=3, value=d["date"].strftime("%a"))
    sm.cell(row=i, column=4, value=SESSION_BY_WEEKDAY[wd])
    sm.cell(row=i, column=5, value=f"=Dashboard!$C$4-(Dashboard!$C$4-Dashboard!$C$5)*B{i}/{DAYS}").number_format = "0.0"
    sm.cell(row=i, column=6, value="=Dashboard!$C$8")
    for col, src in zip("GHIJK", "IJKLM"):
        sm.cell(row=i, column="ABCDEFGHIJK".index(col) + 1,
                value=f"=SUMIFS('Daily Meals'!${src}$2:${src}${LAST_MEAL_ROW},'Daily Meals'!$A$2:$A${LAST_MEAL_ROW},A{i})").number_format = "0"
    lo = max(2, i - 6)
    sm.cell(row=i, column=13, value=f'=IF(COUNT(L{lo}:L{i})=0,"",AVERAGE(L{lo}:L{i}))').number_format = "0.0"
    for col in (12, 14, 15, 16, 17, 18, 19):
        sm.cell(row=i, column=col).fill = INP
    if d["date"] == DIWALI: sm.cell(row=i, column=19, value="Diwali: planned mithai in pre-workout/evening snack. Log it, stay on plan otherwise.")
sm.freeze_panes = "E2"
dv_yn = DataValidation(type="list", formula1='"Y,N"', allow_blank=True); sm.add_data_validation(dv_yn); dv_yn.add(f"Q2:Q{DAYS + 1}")
dv_rd = DataValidation(type="list", formula1='"1,2,3,4,5"', allow_blank=True); sm.add_data_validation(dv_rd); dv_rd.add(f"P2:P{DAYS + 1}")

# chart on dashboard
ch = LineChart(); ch.title = "Weight: actual vs target trajectory"; ch.height = 9; ch.width = 20
ch.y_axis.title = "kg"; ch.x_axis.number_format = "dd-mmm"
for col in (5, 12, 13):
    ch.add_data(Reference(sm, min_col=col, min_row=1, max_row=DAYS + 1), titles_from_data=True)
ch.set_categories(Reference(sm, min_col=1, min_row=2, max_row=DAYS + 1))
ds.add_chart(ch, "E3")

# --- Training Plan
tp = wb.create_sheet("Training Plan")
hdr(tp, 1, ["Session", "Exercise", "Sets", "Reps", "RIR", "Notes"], [10, 40, 6, 9, 7, 80])
r = 2
for s, ex in SESSIONS.items():
    for e in ex:
        for ci, v in enumerate((s, *e), 1): tp.cell(row=r, column=ci, value=v)
        r += 1
    r += 1
tp.cell(row=r, column=1, value="Weekly schedule: Mon Push | Tue Pull | Wed Legs | Thu Rest | Fri Upper | Sat Lower | Sun Rest. Cap each session at ~75 min (rest 2-3 min compounds, ~90 s isolations, superset isolations).").font = Font(bold=True)
r += 2
hdr(tp, r, ["Week", "Dates", "Phase", "Volume", "Intensity guidance", ""])
for wk, phase, vol, rir in BLOCK:
    r += 1
    a = START + dt.timedelta(days=(wk - 1) * 7); b = min(a + dt.timedelta(days=6), START + dt.timedelta(days=DAYS - 1))
    for ci, v in enumerate((wk, f"{a:%d %b} - {b:%d %b}", phase, vol, rir), 1): tp.cell(row=r, column=ci, value=v)
    if phase == "DELOAD":
        for ci in range(1, 6): tp.cell(row=r, column=ci).fill = SUB
tp.cell(row=r + 2, column=1, value="Progression rule: hit the top of the rep range on all sets at target RIR → add the smallest load jump next session. If reps fall 2 sessions in a row → hold the load; 3 in a row or readiness low → trigger recovery protocol.")

# --- Training Log
tl = wb.create_sheet("Training Log")
hdr(tl, 1, ["Date", "Session", "Exercise", "Set #", "Load (kg)", "Reps", "RIR", "Notes"], [12, 10, 36, 7, 10, 7, 7, 40])
for rr in range(2, 400):
    for cc in range(1, 9): tl.cell(row=rr, column=cc).fill = INP
tl.freeze_panes = "A2"

# --- Weekly Check-in
wc = wb.create_sheet("Weekly Check-in")
hdr(wc, 1, ["Week", "Start", "End", "Avg weight (kg)", "Prev avg (kg)", "Change (kg)", "Change %", "Avg steps", "Avg sleep (h)", "Sessions done", "Strength trend",
            "Hunger 1-5", "Adherence %", "Waist (in)", "Status", "Recommended action (one lever only)", "Coach notes"],
    [6, 11, 11, 11, 11, 10, 9, 9, 9, 9, 11, 9, 11, 9, 24, 70, 40])
dvs = DataValidation(type="list", formula1='"Up,Same,Down"', allow_blank=True); wc.add_data_validation(dvs)
for wk in range(1, 14):
    i = wk + 1
    a = START + dt.timedelta(days=(wk - 1) * 7); b = min(a + dt.timedelta(days=6), START + dt.timedelta(days=DAYS - 1))
    wc.cell(row=i, column=1, value=wk)
    wc.cell(row=i, column=2, value=a).number_format = "dd-mmm"; wc.cell(row=i, column=3, value=b).number_format = "dd-mmm"
    rng = lambda col: f"'Daily Summary'!${col}$2:${col}${DAYS + 1}"
    cond = f"{rng('A')},\">=\"&B{i},{rng('A')},\"<=\"&C{i}"
    wc.cell(row=i, column=4, value=f'=IFERROR(AVERAGEIFS({rng("L")},{cond}),"")').number_format = "0.0"
    wc.cell(row=i, column=5, value="=Dashboard!$C$4" if wk == 1 else f"=E{i - 1}*0+D{i - 1}" ).number_format = "0.0"
    wc.cell(row=i, column=6, value=f'=IF(D{i}="","",D{i}-E{i})').number_format = "0.00"
    wc.cell(row=i, column=7, value=f'=IF(D{i}="","",F{i}/E{i})').number_format = "0.0%"
    wc.cell(row=i, column=8, value=f'=IFERROR(AVERAGEIFS({rng("N")},{cond}),"")').number_format = "0"
    wc.cell(row=i, column=9, value=f'=IFERROR(AVERAGEIFS({rng("O")},{cond}),"")').number_format = "0.0"
    wc.cell(row=i, column=10, value=f'=COUNTIFS({rng("Q")},"Y",{cond})')
    dvs.add(f"K{i}")
    for col in (11, 12, 13, 14, 17): wc.cell(row=i, column=col).fill = INP
    wc.cell(row=i, column=15, value=(
        f'=IF(D{i}="","",IF(G{i}>-0.003,"Too slow / stalled",IF(G{i}<-0.01,"Too fast","On track")))'))
    wc.cell(row=i, column=16, value=(
        f'=IF(D{i}="","Enter daily weights first",'
        f'IF(AND(G{i}<-0.01,K{i}="Down"),"Muscle at risk: +150-200 kcal (carbs) and cut volume 20%",'
        f'IF(G{i}<-0.01,"Losing too fast: +150 kcal",'
        f'IF(AND(G{i}>-0.003,M{i}<>"",M{i}<0.85),"Stalled but adherence <85%: fix adherence, NO plan change",'
        f'IF(G{i}>-0.003,"Stalled: -150 kcal OR +1,000 steps/day (pick one)",'
        f'IF(K{i}="Down","Weight OK, strength down: check sleep/readiness; hold calories; trim volume 20%",'
        f'"On track: no change"))))))'))
wc.cell(row=16, column=1, value="Week 1 weight often drops 1-2 kg from water/glycogen — judge the real rate from weeks 2-3. Compare 7-day averages, never single days. Change ONE lever per week.")
# prev avg: carry the last non-blank average forward
for wk in range(2, 14):
    i = wk + 1
    wc.cell(row=i, column=5, value=f'=IF(D{i - 1}="",E{i - 1},D{i - 1})').number_format = "0.0"

# --- Measurements
ms = wb.create_sheet("Measurements")
hdr(ms, 1, ["Date", "Weight (kg)", "Waist (in)", "Chest (in)", "Shoulders (in)", "Arm (in)", "Thigh (in)", "Photos taken?", "Notes"], [12, 11, 10, 10, 13, 9, 10, 13, 40])
for i in range(14):
    d0 = START + dt.timedelta(days=7 * i) if i < 13 else START + dt.timedelta(days=DAYS - 1)
    ms.cell(row=i + 2, column=1, value=d0).number_format = "dd-mmm-yy"
    for cc in range(2, 10): ms.cell(row=i + 2, column=cc).fill = INP
ms["B2"] = args.start_weight; ms["C2"] = 38.5; ms["I2"] = "Baseline; waist is approximate (38-40) — re-measure properly"
ms["A18"] = "Measure on the same morning each week: waist at navel, relaxed, same tape. Photos: front/side/back, same light and time."

# --- Plan & Rules
pr = wb.create_sheet("Plan & Rules")
pr.column_dimensions["A"].width = 120
lines = [
    ("TARGETS", True),
    (f"Calories {args.kcal} kcal/day | Protein ~170-180 g | Fat 55-90 g | Carbs ~230-250 g | Fibre 35-45 g. Weigh raw dal/rice/atta unless noted.", False),
    ("Calories are an estimate (TDEE ~2,900-3,000 assuming ~10k steps + 5 lifts/wk). The weekly check-in recalibrates from your real weight trend.", False),
    ("", False),
    ("DECISION TABLE (weekly, from 7-day avg weight)", True),
    ("Loss 0.5-0.9 kg/week (0.45-0.9%) AND strength Same/Up → hold everything.", False),
    ("Loss < 0.3%/week for 2 weeks AND adherence ≥85% → -150 kcal OR +1,000 steps (one lever).", False),
    ("Loss > 1%/week OR strength down 2 sessions → +150-200 kcal from carbs, volume -20%.", False),
    ("Hunger ≥4/5 and sleep < 6.5 h for a week → fix sleep first; consider a 5-7 day maintenance week.", False),
    ("", False),
    ("SUPPLEMENTS", True),
    ("Whey isolate: 1-2 scoops/day (already counted). Use isolate only — drop the concentrate if skin is flaring (whey/dairy–acne link is observational; test over 4 weeks).", False),
    ("Vitamin D3: 60,000 IU once weekly with your largest fat-containing meal, as prescribed by your nutritionist; typically 8 weeks then re-test 25(OH)D. Confirm duration with your doctor.", False),
    ("Creatine monohydrate 3-5 g/day (optional, well-supported for strength/muscle retention; expect +0.5-1 kg water in week 1-2).", False),
    ("Consider testing B12 (vegetarian) with your next blood panel.", False),
    ("", False),
    ("EVENTS", True),
    ("Diwali (8 Nov, Sunday rest day): the plan includes 2 small mithai in the evening snack. Anything beyond: log it, no 'make up' fasting, just return to plan next meal.", False),
    ("Bangalore trip: tell the coach the dates. Default rules — protein first (paneer tikka, egg bhurji/omelette, dahi, dal), skip sugary drinks, hit steps by walking, keep 3 workouts or none.", False),
]
for i, (t, b) in enumerate(lines, 1):
    c = pr.cell(row=i, column=1, value=t); c.font = Font(bold=b); c.alignment = Alignment(wrap_text=True)

os.makedirs(os.path.dirname(os.path.abspath(args.out)), exist_ok=True)
wb.save(args.out)
ks = [d["tot"][0] for d in days]; ps = [d["tot"][1] for d in days]; fs = [d["tot"][3] for d in days]
print(f"Saved {args.out}: {len(days)} days | kcal {min(ks):.0f}-{max(ks):.0f} | protein {min(ps):.0f}-{max(ps):.0f} g | fat {min(fs):.0f}-{max(fs):.0f} g")
